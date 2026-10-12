#!/usr/bin/env python3
"""Post-process the html2pptx deck: rotate the vertical METHOD rail label.

html2pptx emits the rotated side-title as an upright, narrow text box, so
PowerPoint shows "ME / TH / OD" stacked. This swaps the box's width/height,
centres it on the same spot and applies a 270-degree rotation so it reads
bottom-to-top like the HTML poster.

Usage: python assets/meta/fix_pptx.py <in.pptx> <out.pptx>
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = Path(sys.argv[1]), Path(sys.argv[2])
prs = Presentation(str(src))
slide = prs.slides[0]
fixed = 0
for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    if sh.text_frame.text.strip() != "METHOD":
        continue
    if sh.width >= sh.height:
        continue  # already landscape-shaped / rotated
    cx, cy = sh.left + sh.width // 2, sh.top + sh.height // 2
    new_w, new_h = sh.height, sh.width
    sh.width, sh.height = new_w, new_h
    sh.left, sh.top = cx - new_w // 2, cy - new_h // 2
    sh.rotation = 270
    tf = sh.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for para in tf.paragraphs:
        para.alignment = PP_ALIGN.CENTER
    fixed += 1
prs.save(str(dst))
print(f"rotated {fixed} METHOD label(s) -> {dst}")
